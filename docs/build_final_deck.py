#!/usr/bin/env python3.12
"""Generate the final-demo presentation (docs/FinalPresentation.pptx).

Matches the dark theme of the mid-demo deck: near-black background, a green/blue
gradient strip on top, two-column rounded "cards", green/blue/orange section
headers, monospace code blocks, and "X / N" page numbers.

Focus: the four extensions delivered since the mid-demo —
Shadow Mapping, PBR + IBL, Inverse Kinematics, Animation State Machine.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ── Palette ──────────────────────────────────────────────────────────────────
BG       = RGBColor(0x0B, 0x0E, 0x12)
CARD     = RGBColor(0x16, 0x1B, 0x22)
CARD_BRD = RGBColor(0x2A, 0x32, 0x3C)
CODE_BG  = RGBColor(0x0D, 0x11, 0x17)
GREEN    = RGBColor(0x2E, 0xE6, 0x8A)   # bright accent green
GREEN_H  = RGBColor(0x35, 0xD3, 0x7A)   # header green
BLUE     = RGBColor(0x3F, 0xA9, 0xF5)
ORANGE   = RGBColor(0xFF, 0x7A, 0x3C)
WHITE    = RGBColor(0xEC, 0xEF, 0xF1)
GREY     = RGBColor(0xA8, 0xB2, 0xBD)
MUTED    = RGBColor(0x76, 0x82, 0x8E)
CODE_FG  = RGBColor(0xC8, 0xD3, 0xDE)

FONT      = "Calibri"
MONO      = "Consolas"

EMU_W, EMU_H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width  = Emu(12192000)   # exact 13.333in (16:9)
prs.slide_height = Emu(6858000)    # exact 7.5in
# python-pptx's default template leaves type="screen4x3" on <p:sldSz>, which now
# mismatches the 16:9 size and makes Keynote reject the file. Drop the type attr.
_sldSz = prs._element.find(qn('p:sldSz'))
if _sldSz is not None and 'type' in _sldSz.attrib:
    del _sldSz.attrib['type']
BLANK = prs.slide_layouts[6]

TOTAL = 13  # filled at end into page numbers


def _solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _noautosize(tf):
    # prevent pptx from auto-growing/shrinking text frames
    tf.word_wrap = True
    el = tf._txBody
    bodyPr = el.find(qn('a:bodyPr'))
    for tag in ('a:spAutoFit', 'a:normAutofit'):
        e = bodyPr.find(qn(tag))
        if e is not None:
            bodyPr.remove(e)


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, EMU_W, EMU_H)
    _solid(bg, BG)
    bg.shadow.inherit = False
    # top gradient strip: green left half -> blue right half
    half = Emu(int(EMU_W) // 2)
    g = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, half, Inches(0.09))
    _solid(g, GREEN); g.shadow.inherit = False
    b = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, half, 0, half, Inches(0.09))
    _solid(b, BLUE); b.shadow.inherit = False
    return s


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
         space_after=6, line_spacing=1.0):
    """runs: list of paragraphs; each paragraph is list of (txt, size, color, bold, font)."""
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    _noautosize(tf)
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        for (txt, size, color, bold, font) in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.color.rgb = color
            r.font.bold = bold; r.font.name = font
    return tb


def title_bar(s, title, subtitle=None):
    runs = [[(title, 30, GREEN_H, True, FONT)]]
    text(s, Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.8), runs)
    if subtitle:
        text(s, Inches(0.57), Inches(0.95), Inches(12.0), Inches(0.4),
             [[(subtitle, 12.5, MUTED, False, FONT)]])


def page_num(s, n):
    text(s, Inches(11.7), Inches(7.02), Inches(1.4), Inches(0.4),
         [[(f"{n} / {TOTAL}", 11, MUTED, False, FONT)]], align=PP_ALIGN.RIGHT)


def card(s, x, y, w, h, accent):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    c.adjustments[0] = 0.04
    c.fill.solid(); c.fill.fore_color.rgb = CARD
    c.line.color.rgb = CARD_BRD; c.line.width = Pt(0.75)
    c.shadow.inherit = False
    bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y,
                             Inches(0.07), h)
    bar.adjustments[0] = 0.5
    _solid(bar, accent); bar.shadow.inherit = False
    return c


def card_header(s, x, y, w, txt, color):
    text(s, x + Inches(0.28), y + Inches(0.14), w - Inches(0.4), Inches(0.4),
         [[(txt, 16, color, True, FONT)]])


def bullets(s, x, y, w, h, items, size=13.5, color=WHITE, gap=7, bullet="•",
            bcolor=None):
    bcolor = bcolor or GREEN
    runs = []
    for it in items:
        if isinstance(it, tuple):
            txt, c = it
        else:
            txt, c = it, color
        runs.append([(f"{bullet}  ", size, bcolor, True, FONT),
                     (txt, size, c, False, FONT)])
    text(s, x, y, w, h, runs, space_after=gap, line_spacing=1.04)


def code_block(s, x, y, w, h, lines, size=11.5):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    c.adjustments[0] = 0.06
    c.fill.solid(); c.fill.fore_color.rgb = CODE_BG
    c.line.color.rgb = CARD_BRD; c.line.width = Pt(0.5)
    c.shadow.inherit = False
    runs = [[(ln, size, CODE_FG, False, MONO)] for ln in lines]
    text(s, x + Inches(0.2), y + Inches(0.13), w - Inches(0.35),
         h - Inches(0.25), runs, space_after=2, line_spacing=1.05)


def chip(s, x, y, w, h, big, small, color):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    c.adjustments[0] = 0.12
    c.fill.solid(); c.fill.fore_color.rgb = CARD
    c.line.color.rgb = color; c.line.width = Pt(1.25)
    c.shadow.inherit = False
    text(s, x, y + Inches(0.16), w, Inches(0.5),
         [[(big, 22, color, True, FONT)]], align=PP_ALIGN.CENTER)
    text(s, x, y + Inches(0.66), w, Inches(0.35),
         [[(small, 11, MUTED, False, FONT)]], align=PP_ALIGN.CENTER)


# Common geometry for two-column slides
LX, RX = Inches(0.55), Inches(6.95)
CW = Inches(5.85)
CY = Inches(1.35)
CH = Inches(5.35)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 1 — Title
# ══════════════════════════════════════════════════════════════════════════════
s = slide()
# left accent bar
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(1.25),
                         Inches(0.08), Inches(2.6))
_solid(bar, GREEN); bar.shadow.inherit = False
text(s, Inches(0.85), Inches(1.15), Inches(11.5), Inches(2.2),
     [[("Real-Time Skeletal", 46, WHITE, True, FONT)],
      [("Animation Viewer", 46, WHITE, True, FONT)]], line_spacing=1.0)
text(s, Inches(0.9), Inches(3.55), Inches(11), Inches(0.5),
     [[("Final Demo — Rendering & Animation Extensions", 19, GREEN, False, FONT)]])
ln = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(4.25),
                        Inches(5.6), Pt(1.2))
_solid(ln, CARD_BRD); ln.shadow.inherit = False
text(s, Inches(0.9), Inches(4.45), Inches(11), Inches(0.4),
     [[("C++17  ·  OpenGL 4.1 Core  ·  glTF 2.0", 14, MUTED, False, FONT)]])
text(s, Inches(0.9), Inches(5.05), Inches(11), Inches(0.4),
     [[("Mehmet Yiğit Tutaş  ·  Ozan Özak", 15, WHITE, False, FONT)]])
text(s, Inches(0.9), Inches(5.6), Inches(11), Inches(0.4),
     [[("Computer Graphics Course Project", 13, MUTED, False, FONT)]])
page_num(s, 1)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 2 — What's New (brief recap + the four extensions)
# ══════════════════════════════════════════════════════════════════════════════
s = slide()
title_bar(s, "What's New Since the Mid-Demo", "Recap · Delivered Extensions")
card(s, LX, CY, CW, CH, GREEN)
card_header(s, LX, CY, CW, "Recap — Base Viewer (Mid-Demo)", GREEN_H)
bullets(s, LX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.55), Inches(3),
        ["glTF 2.0 loading — geometry, skeleton, animation, textures",
         "Forward Kinematics: global[i] = global[parent] × local[i]",
         "GPU Linear Blend Skinning (4 bone influences / vertex)",
         "TRS-level blending with SLERP · ImGui debug panel",
         "60 FPS on CesiumMan & RiggedFigure"])
text(s, LX + Inches(0.3), CY + Inches(3.5), CW - Inches(0.6), Inches(1.6),
     [[("The base pipeline is the foundation; the four extensions "
        "below build directly on its FK + skinning core.", 12.5, MUTED, False, FONT)]],
     line_spacing=1.1)

card(s, RX, CY, CW, CH, BLUE)
card_header(s, RX, CY, CW, "Four New Extensions", BLUE)
items = [
    ("Shadow Mapping", "two-pass depth · PCF soft shadows", GREEN),
    ("PBR + Image-Based Lighting", "metallic-roughness BRDF · split-sum IBL", BLUE),
    ("Inverse Kinematics", "two-bone analytic solver", ORANGE),
    ("Animation State Machine", "keyboard Walk ⇄ Run crossfade", GREEN),
]
yy = CY + Inches(0.85)
for i, (t, d, col) in enumerate(items):
    num = s.shapes.add_shape(MSO_SHAPE.OVAL, RX + Inches(0.3), yy,
                             Inches(0.5), Inches(0.5))
    _solid(num, col); num.shadow.inherit = False
    text(s, RX + Inches(0.3), yy + Inches(0.04), Inches(0.5), Inches(0.4),
         [[(str(i + 1), 16, BG, True, FONT)]], align=PP_ALIGN.CENTER)
    text(s, RX + Inches(0.98), yy - Inches(0.03), CW - Inches(1.2), Inches(0.5),
         [[(t, 15.5, WHITE, True, FONT)]])
    text(s, RX + Inches(0.98), yy + Inches(0.34), CW - Inches(1.2), Inches(0.4),
         [[(d, 12, MUTED, False, FONT)]])
    yy += Inches(1.05)
page_num(s, 2)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 3 — Shadow Mapping: concept
# ══════════════════════════════════════════════════════════════════════════════
s = slide()
title_bar(s, "Shadow Mapping — The Idea",
          "Same question as a shadow ray, answered ahead of time")

# Left card: the one question + ray/path tracing
card(s, LX, CY, CW, CH, GREEN)
card_header(s, LX, CY, CW, "The One Question", GREEN_H)
bullets(s, LX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.55), Inches(1.7),
        ["A spot is in shadow when something stands between it and the light",
         "So the only question is: can this spot see the light, or not?"],
        gap=10)
# two "direct answer" methods
text(s, LX + Inches(0.3), CY + Inches(2.35), CW - Inches(0.6), Inches(1.3),
     [[("Ray tracing", 14, ORANGE, True, FONT),
       ("  — send one ray straight to the light;", 12.5, WHITE, False, FONT)],
      [("is anything in the way?", 12.5, GREY, False, FONT)]],
     line_spacing=1.1, space_after=4)
text(s, LX + Inches(0.3), CY + Inches(3.45), CW - Inches(0.6), Inches(1.4),
     [[("Path tracing", 14, BLUE, True, FONT),
       ("  — many bouncing rays; gorgeous,", 12.5, WHITE, False, FONT)],
      [("but far too slow for real time.", 12.5, GREY, False, FONT)]],
     line_spacing=1.1)

# Right card: the shortcut, two steps
card(s, RX, CY, CW, CH, ORANGE)
card_header(s, RX, CY, CW, "Shadow Mapping — Our Shortcut", ORANGE)
yy = CY + Inches(0.85)
for n, txt in [("1", "The light takes a “picture” — for every direction, "
                     "how far away is the nearest thing?"),
               ("2", "While drawing, each spot asks: am I farther away than the "
                     "nearest thing the light saw?")]:
    num = s.shapes.add_shape(MSO_SHAPE.OVAL, RX + Inches(0.3), yy,
                             Inches(0.5), Inches(0.5))
    _solid(num, ORANGE); num.shadow.inherit = False
    text(s, RX + Inches(0.3), yy + Inches(0.04), Inches(0.5), Inches(0.4),
         [[(n, 16, BG, True, FONT)]], align=PP_ALIGN.CENTER)
    text(s, RX + Inches(0.98), yy - Inches(0.02), CW - Inches(1.25), Inches(1.1),
         [[(txt, 13, WHITE, False, FONT)]], line_spacing=1.08)
    yy += Inches(1.35)
# result line
rb = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, RX + Inches(0.3),
                        yy + Inches(0.05), CW - Inches(0.6), Inches(0.7))
rb.adjustments[0] = 0.2
rb.fill.solid(); rb.fill.fore_color.rgb = CARD; rb.line.color.rgb = ORANGE
rb.line.width = Pt(1); rb.shadow.inherit = False
text(s, RX + Inches(0.45), yy + Inches(0.18), CW - Inches(0.9), Inches(0.5),
     [[("Farther → something is blocking it → in shadow", 13, WHITE, True, FONT)]])
text(s, RX + Inches(0.3), yy + Inches(0.95), CW - Inches(0.6), Inches(1),
     [[("The same “can I see the light?” test as a shadow ray — just "
        "done once, ahead of time, instead of one ray at a time.",
        11.5, MUTED, False, FONT)]], line_spacing=1.12)
page_num(s, 3)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 4 — Shadow Mapping: soft edges (plain-language, matches the script)
# ══════════════════════════════════════════════════════════════════════════════
s = slide()
title_bar(s, "Shadow Mapping — Soft Edges",
          "Why edges look jagged — and the simple fix")

# Left card: the problem
card(s, LX, CY, CW, CH, GREEN)
card_header(s, LX, CY, CW, "The Problem — Jagged Edges", GREEN_H)
bullets(s, LX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.55), Inches(3.2),
        ["The light's picture is fairly low-resolution",
         "Each spot gets a strict yes / no answer — fully shadowed or "
         "fully lit, nothing in between",
         "So the shadow boundary snaps to those pixels — a blocky, "
         "staircased edge"], gap=12)
text(s, LX + Inches(0.3), CY + Inches(3.95), CW - Inches(0.6), Inches(1.1),
     [[("Hard edge", 13, GREEN, True, FONT)],
      [("= asking one person \"is it raining here?\" → just yes or no.",
        12.5, GREY, False, FONT)]], line_spacing=1.15)

# Right card: the fix
card(s, RX, CY, CW, CH, BLUE)
card_header(s, RX, CY, CW, "The Fix — Average the Neighbours", BLUE)
bullets(s, RX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.55), Inches(2.0),
        ["Instead of checking one point, check a few neighbouring points",
         "Average their yes / no answers → partial shadow right at the edge",
         "That partial value becomes a soft, gradual fade, not a hard line"],
        bcolor=BLUE, gap=10)

# Small intuitive 3x3 'vote' grid
gx, gy = RX + Inches(0.45), CY + Inches(3.05)
cell = Inches(0.42); padc = Inches(0.12)
pattern = [1, 1, 0,
           1, 1, 0,
           1, 0, 0]   # 5 of 9 shadowed
for i, v in enumerate(pattern):
    r, c = divmod(i, 3)
    cx = Emu(int(gx) + c * (int(cell) + int(padc)))
    cy = Emu(int(gy) + r * (int(cell) + int(padc)))
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, cx, cy, cell, cell)
    dot.shadow.inherit = False
    if v:
        dot.fill.solid(); dot.fill.fore_color.rgb = RGBColor(0x2A, 0x32, 0x3C)
        dot.line.color.rgb = MUTED
    else:
        dot.fill.solid(); dot.fill.fore_color.rgb = RGBColor(0xF2, 0xE9, 0xC0)
        dot.line.color.rgb = RGBColor(0xC9, 0xB8, 0x6A)
    dot.line.width = Pt(1)
text(s, gx + Inches(1.9), gy + Inches(0.2), CW - Inches(2.6), Inches(1.6),
     [[("5 of 9 shadowed", 14, WHITE, True, FONT)],
      [("→ about 55% dark —", 13, BLUE, False, FONT)],
      [("a soft middle value", 13, BLUE, False, FONT)]], line_spacing=1.1)
text(s, RX + Inches(0.3), CY + Inches(4.75), CW - Inches(0.6), Inches(0.5),
     [[("A cheap stand-in for the naturally soft shadows a real light gives.",
        11.5, MUTED, False, FONT)]])
page_num(s, 4)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 5 — PBR: Cook-Torrance BRDF
# ══════════════════════════════════════════════════════════════════════════════
s = slide()
title_bar(s, "Physically Based Rendering",
          "Three simple values, real-world light behaviour")

# Left card: the three values
card(s, LX, CY, CW, CH, GREEN)
card_header(s, LX, CY, CW, "Describe Any Surface With 3 Values", GREEN_H)
vals = [("Base Colour", "the basic colour of the surface"),
        ("Metallic", "is it metal or not?  0 = plastic / wood,  1 = metal"),
        ("Roughness", "smooth = sharp mirror,  rough = soft matte")]
yy = CY + Inches(0.8)
for name, desc in vals:
    text(s, LX + Inches(0.3), yy, CW - Inches(0.6), Inches(0.45),
         [[(name, 15, GREEN, True, FONT)]])
    text(s, LX + Inches(0.3), yy + Inches(0.36), CW - Inches(0.6), Inches(0.5),
         [[(desc, 12.5, GREY, False, FONT)]], line_spacing=1.05)
    yy += Inches(0.95)
# demo callout
db = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, LX + Inches(0.3),
                        yy + Inches(0.05), CW - Inches(0.6), Inches(1.15))
db.adjustments[0] = 0.1
db.fill.solid(); db.fill.fore_color.rgb = CARD; db.line.color.rgb = GREEN
db.line.width = Pt(1); db.shadow.inherit = False
text(s, LX + Inches(0.48), yy + Inches(0.2), CW - Inches(0.9), Inches(0.9),
     [[("Live demo", 13, GREEN, True, FONT)],
      [("Drag Metallic and Roughness — the same model goes from matte "
        "plastic to a chrome mirror.", 12, GREY, False, FONT)]],
     line_spacing=1.1)

# Right card: how it lights each pixel
card(s, RX, CY, CW, CH, BLUE)
card_header(s, RX, CY, CW, "How It Lights Each Pixel", BLUE)
bullets(s, RX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.55), Inches(2.6),
        ["Add two things: the soft, even base colour, and the shiny "
         "highlights and reflections on top",
         "Roughness decides how sharp those highlights are",
         "Built-in real-world effect: surfaces get more mirror-like edge-on — "
         "a tabletop is dull from above, reflective when you look across it"],
        bcolor=BLUE, gap=11)
rb = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, RX + Inches(0.3),
                        CY + Inches(3.95), CW - Inches(0.6), Inches(0.95))
rb.adjustments[0] = 0.14
rb.fill.solid(); rb.fill.fore_color.rgb = CARD; rb.line.color.rgb = GREEN
rb.line.width = Pt(1); rb.shadow.inherit = False
text(s, RX + Inches(0.48), CY + Inches(4.18), CW - Inches(0.9), Inches(0.6),
     [[("It follows real physics — so it looks correct under any lighting.",
        13, WHITE, True, FONT)]], line_spacing=1.1)
page_num(s, 5)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 6 — IBL split-sum
# ══════════════════════════════════════════════════════════════════════════════
s = slide()
title_bar(s, "Image-Based Lighting",
          "Light the model with its whole surroundings")

# Simplified 3-step flow (plain language)
steps = [("Build a sky", "in code", ORANGE),
         ("Pre-bake into", "small images", BLUE),
         ("Look up while", "rendering", GREEN)]
sw = Inches(3.6); sy = Inches(1.35); sh = Inches(0.9)
for i, (l1, l2, col) in enumerate(steps):
    bx = Emu(int(LX) + i * int(Inches(4.15)))
    bb = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, sy, sw, sh)
    bb.adjustments[0] = 0.12
    bb.fill.solid(); bb.fill.fore_color.rgb = CARD
    bb.line.color.rgb = col; bb.line.width = Pt(1.25); bb.shadow.inherit = False
    text(s, bx, sy + Inches(0.14), sw, sh,
         [[(l1, 14, WHITE, True, FONT)], [(l2, 14, WHITE, True, FONT)]],
         align=PP_ALIGN.CENTER, space_after=1)
    if i < 2:
        ax = Emu(int(bx) + int(sw) + Inches(0.08))
        text(s, ax, sy + Inches(0.18), Inches(0.45), Inches(0.5),
             [[("→", 22, MUTED, True, FONT)]], align=PP_ALIGN.CENTER)

cy2 = Inches(2.5); ch2 = Inches(3.6)
# Left card: the idea
card(s, LX, cy2, CW, ch2, GREEN)
card_header(s, LX, cy2, CW, "The Idea", GREEN_H)
bullets(s, LX + Inches(0.3), cy2 + Inches(0.7), CW - Inches(0.55), Inches(2.7),
        ["Light the character with its whole environment, not just one lamp",
         "Metal reflects the sky; the shaded side picks up colour from "
         "around it",
         "Replaces the old flat, lifeless constant ambient"], gap=12)

# Right card: the trick
card(s, RX, cy2, CW, ch2, BLUE)
card_header(s, RX, cy2, CW, "The Trick — Bake Once, Look Up", BLUE)
bullets(s, RX + Inches(0.3), cy2 + Inches(0.7), CW - Inches(0.55), Inches(2.7),
        ["Adding light from every direction, for every pixel, is too heavy "
         "to do live",
         "So we do it once, up front, and save it as a few small images to "
         "look up",
         "A soft, blurry sky → gentle fill light",
         "Sharper-to-blurrier skies → reflections: smooth = mirror, "
         "rough = hazy"], size=12.5, bcolor=BLUE, gap=9)

# Full-width footer: the path-tracing connection
fb = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, LX, Inches(6.25),
                        Inches(12.25), Inches(0.62))
fb.adjustments[0] = 0.3
fb.fill.solid(); fb.fill.fore_color.rgb = CARD; fb.line.color.rgb = ORANGE
fb.line.width = Pt(1); fb.shadow.inherit = False
text(s, LX + Inches(0.3), Inches(6.38), Inches(11.7), Inches(0.4),
     [[("Same idea as our shadows — a fast stand-in for the all-around "
        "lighting path tracing computes properly, but too slowly for real "
        "time.", 12.5, WHITE, False, FONT)]], align=PP_ALIGN.CENTER)
page_num(s, 6)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 7 — IK concept
# ══════════════════════════════════════════════════════════════════════════════
s = slide()
title_bar(s, "Inverse Kinematics — The Inverse Problem",
          "Position → joint angles (FK run backwards)")
card(s, LX, CY, CW, CH, GREEN)
card_header(s, LX, CY, CW, "Forward vs. Inverse", GREEN_H)
bullets(s, LX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.55), Inches(2.3),
        [("Forward (FK): joint angles → end position. This is the existing "
          "animation pipeline.", WHITE),
         ("Inverse (IK): you give a target position; the solver finds the "
          "joint angles that reach it.", GREEN)], gap=10)
code_block(s, LX + Inches(0.3), CY + Inches(2.55), CW - Inches(0.6), Inches(1.5),
           ["FK:  angles  ──►  position",
            "IK:  position ──►  angles",
            "",
            "Chain:  root → mid → end"])
text(s, LX + Inches(0.3), CY + Inches(4.25), CW - Inches(0.6), Inches(0.8),
     [[("Two-bone analytic IK — closed-form, no iteration, runs every frame.",
        12, MUTED, False, FONT)]], line_spacing=1.1)

card(s, RX, CY, CW, CH, ORANGE)
card_header(s, RX, CY, CW, "How It Solves (Law of Cosines)", ORANGE)
bullets(s, RX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.55), Inches(3.4),
        ["Three joints: root (hip/shoulder), mid (knee/elbow), end (foot/hand)",
         "Bone lengths are fixed; only the distance to the target changes",
         "Law of cosines gives the interior angle at the root and the mid joint",
         "An 'aim' rotation then points the whole chain at the target",
         "Target beyond reach ⇒ the limb straightens (clamped)"],
        bcolor=ORANGE, gap=9)
page_num(s, 7)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 8 — IK integration
# ══════════════════════════════════════════════════════════════════════════════
s = slide()
title_bar(s, "Inverse Kinematics — Integration",
          "Where IK hooks into the per-frame pose pipeline")
card(s, LX, CY, CW, CH, GREEN)
card_header(s, LX, CY, CW, "Inserted After FK, Before Skinning", GREEN_H)
code_block(s, LX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.6), Inches(2.3),
           ["computeLocal();   // sample clips",
            "computeGlobal();  // FK pass",
            "",
            "if (ikEnabled)",
            "  solveTwoBoneIK();",
            "  computeGlobal();  // re-run FK",
            "",
            "computeSkinning();"])
bullets(s, LX + Inches(0.3), CY + Inches(3.2), CW - Inches(0.55), Inches(1.8),
        ["Solver edits the local rotations of root + mid joints",
         "World-space correction → parent-frame local rotation",
         "Re-running FK propagates to the end joint & descendants"],
        size=12.5, gap=6)

card(s, RX, CY, CW, CH, BLUE)
card_header(s, RX, CY, CW, "Interactive Demo Controls", BLUE)
bullets(s, RX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.55), Inches(3.4),
        [("User-selectable end effector — pick any bone from a dropdown; "
          "root & mid are derived from its parents", WHITE),
         ("Yellow cross marker shows the live target", WHITE),
         ("Drag target X/Y/Z — the knee/elbow bends to reach it", WHITE),
         ("Toggle with the I key or the panel checkbox", WHITE),
         ("Composes on top of the playing animation — the body keeps "
          "walking while the limb tracks", GREEN)], bcolor=BLUE, gap=9)
page_num(s, 8)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 9 — State machine design
# ══════════════════════════════════════════════════════════════════════════════
s = slide()
title_bar(s, "Animation State Machine", "Keyboard-triggered Walk ⇄ Run · real-clip crossfade")
card(s, LX, CY, CW, CH, GREEN)
card_header(s, LX, CY, CW, "States & Transitions", GREEN_H)
bullets(s, LX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.55), Inches(2.2),
        ["Keys 1 / 2 (or panel buttons) request Walk / Run",
         "Each request maps to a real clip and starts a 0.3 s crossfade",
         "Reuses the mid-demo TRS-level SLERP blend — no pops or shearing"])
code_block(s, LX + Inches(0.3), CY + Inches(2.75), CW - Inches(0.6), Inches(1.6),
           ["void request(State s) {",
            "  int clip = (s==Run)?run:walk;",
            "  a.blendTo(clip, 0.3f);",
            "}"])
text(s, LX + Inches(0.3), CY + Inches(4.5), CW - Inches(0.6), Inches(0.6),
     [[("Single-clip models (CesiumMan / RiggedFigure) fall back to "
        "speed-scaling: Run = 2.2× Walk.", 12, MUTED, False, FONT)]],
     line_spacing=1.1)

card(s, RX, CY, CW, CH, BLUE)
card_header(s, RX, CY, CW, "New Model: Khronos Fox", BLUE)
bullets(s, RX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.55), Inches(2.4),
        [("24 bones · 3 real clips: Survey / Walk / Run (CC0)", WHITE),
         ("Provides genuinely distinct animations to transition between", WHITE),
         ("Demonstrates true clip-to-clip blending, not just speed change",
          GREEN)], bcolor=BLUE, gap=9)
# honest note about Jump
nb = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, RX + Inches(0.3),
                        CY + Inches(3.25), CW - Inches(0.6), Inches(1.75))
nb.adjustments[0] = 0.08
nb.fill.solid(); nb.fill.fore_color.rgb = CARD
nb.line.color.rgb = ORANGE; nb.line.width = Pt(1); nb.shadow.inherit = False
text(s, RX + Inches(0.5), CY + Inches(3.4), CW - Inches(0.95), Inches(1.5),
     [[("Design decision — Jump dropped", 13, ORANGE, True, FONT)],
      [("No authored jump clip exists for these rigs; rigidly translating "
        "the whole model upward looked artificial, so locomotion stays "
        "Walk ⇄ Run.", 11.5, GREY, False, FONT)]], line_spacing=1.08)
page_num(s, 9)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 10 — State machine architecture
# ══════════════════════════════════════════════════════════════════════════════
s = slide()
title_bar(s, "State Machine — Architecture", "A thin controller over the existing blend API")
card(s, LX, CY, CW, CH, GREEN)
card_header(s, LX, CY, CW, "Additive Design", GREEN_H)
bullets(s, LX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.55), Inches(3.4),
        ["New AnimStateMachine class owns the current state",
         "configureForModel() maps states → clip indices by clip name "
         "(case-insensitive); falls back to speed-scaling when a model has "
         "one clip",
         "request() drives transitions through the existing "
         "Animator::blendTo() — no change to the core pipeline",
         "Keyboard handler & ImGui panel call request()"])

card(s, RX, CY, CW, CH, BLUE)
card_header(s, RX, CY, CW, "Data Flow", BLUE)
code_block(s, RX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.6), Inches(3.0),
           ["key '2'  ──►  request(Run)",
            "                  │",
            "                  ▼",
            "      AnimStateMachine",
            "        maps Run → clipIdx",
            "                  │",
            "                  ▼",
            "   Animator::blendTo(clip,",
            "                0.3s)",
            "                  │",
            "                  ▼",
            "   TRS SLERP crossfade"])
text(s, RX + Inches(0.3), CY + Inches(3.95), CW - Inches(0.6), Inches(1),
     [[("The state machine is pure orchestration — all the heavy lifting "
        "is the blend system already validated in the mid-demo.",
        12, MUTED, False, FONT)]], line_spacing=1.1)
page_num(s, 10)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 11 — Implementation challenges (new)
# ══════════════════════════════════════════════════════════════════════════════
s = slide()
title_bar(s, "Implementation Challenges", "New issues surfaced by the extensions")
chals = [
    ("Fox mesh invisible — non-indexed geometry",
     "The Fox primitive has no index buffer. Synthesised a sequential "
     "index list (0,1,2,…) so glDrawElements works."),
    ("Flat shading — Fox has no normals",
     "The asset ships without a NORMAL attribute. Generated smooth "
     "per-vertex normals by accumulating face normals."),
    ("Mesh collapsed during Walk — sparse channels",
     "Fox clips animate rotation only on most joints. Missing T/S tracks "
     "now fall back to the bone's BIND transform instead of zero."),
    ("IBL on GL 4.1 — no compute shaders / SSBO",
     "All environment maps baked offline via render-to-cubemap FBOs and "
     "fragment-shader importance sampling."),
    ("IK — world vs. local rotation",
     "Solver computes a world-space correction; converted into the bone's "
     "parent frame before writing the local rotation."),
]
cyc = Inches(1.3); chh = Inches(1.03); gap = Inches(0.1)
for i, (h, d) in enumerate(chals):
    yy = Emu(int(cyc) + i * (int(chh) + int(gap)))
    c = card(s, LX, yy, Inches(12.25), chh, ORANGE)
    text(s, LX + Inches(0.3), yy + Inches(0.1), Inches(11.7), Inches(0.4),
         [[(h, 14, ORANGE, True, FONT)]])
    text(s, LX + Inches(0.3), yy + Inches(0.46), Inches(11.6), Inches(0.5),
         [[(d, 12, GREY, False, FONT)]], line_spacing=1.0)
page_num(s, 11)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 12 — Demo guide / results
# ══════════════════════════════════════════════════════════════════════════════
s = slide()
title_bar(s, "Live Demo Guide", "What to show · controls for the new features")
card(s, LX, CY, CW, CH, GREEN)
card_header(s, LX, CY, CW, "Demo Flow", GREEN_H)
bullets(s, LX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.55), Inches(4.2),
        [("Shadows — toggle ground/shadows; drag the light azimuth/elevation "
          "to swing the soft shadow", WHITE),
         ("PBR — open Override Material; slide Metallic→1, Roughness→0 for a "
          "chrome mirror of the sky", WHITE),
         ("IBL — toggle IBL Ambient & Skybox to show environment lighting", WHITE),
         ("Switch to Fox (M) — press 2/1 to crossfade Run ⇄ Walk", WHITE),
         ("IK — pick a leg/arm bone, press I, drag the target marker", WHITE)],
        size=13, gap=11)

card(s, RX, CY, CW, CH, BLUE)
card_header(s, RX, CY, CW, "Key Bindings (new)", BLUE)
code_block(s, RX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.6), Inches(2.2),
           ["1 / 2    Walk / Run state",
            "I        Toggle IK",
            "M        Switch model",
            "B        Bone overlay",
            "LMB/RMB  Orbit / Pan",
            "Scroll   Zoom"])
chip(s, RX + Inches(0.3), CY + Inches(3.2), Inches(1.7), Inches(1.15),
     "60", "FPS sustained", GREEN)
chip(s, RX + Inches(2.15), CY + Inches(3.2), Inches(1.7), Inches(1.15),
     "3", "models", BLUE)
chip(s, RX + Inches(4.0), CY + Inches(3.2), Inches(1.7), Inches(1.15),
     "4", "extensions", ORANGE)
page_num(s, 12)


# ══════════════════════════════════════════════════════════════════════════════
# SLIDE 13 — Conclusion
# ══════════════════════════════════════════════════════════════════════════════
s = slide()
title_bar(s, "Conclusion", "All four mid-demo 'future extensions' delivered")
card(s, LX, CY, CW, CH, GREEN)
card_header(s, LX, CY, CW, "Delivered", GREEN_H)
bullets(s, LX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.55), Inches(4),
        ["Shadow mapping — two-pass depth, PCF, slope-scaled bias",
         "PBR — Cook-Torrance metallic-roughness BRDF",
         "Image-based lighting — split-sum IBL + procedural skybox",
         "Inverse kinematics — interactive two-bone solver",
         "Animation state machine — Walk ⇄ Run real-clip crossfade",
         "Added the Khronos Fox (multi-clip) reference model"])

card(s, RX, CY, CW, CH, BLUE)
card_header(s, RX, CY, CW, "Takeaways & Future Work", BLUE)
bullets(s, RX + Inches(0.3), CY + Inches(0.7), CW - Inches(0.55), Inches(2.4),
        [("Extensions build additively on the FK + skinning core — the base "
          "pipeline was the right foundation", WHITE),
         ("Real-time techniques mirror offline ones (shadow map ↔ shadow ray, "
          "split-sum ↔ path-traced IBL)", WHITE)], bcolor=BLUE, gap=10)
text(s, RX + Inches(0.3), CY + Inches(2.9), CW - Inches(0.6), Inches(0.4),
     [[("Future work", 14, GREEN, True, FONT)]])
bullets(s, RX + Inches(0.3), CY + Inches(3.4), CW - Inches(0.55), Inches(1.6),
        ["CCD / FABRIK multi-bone IK chains",
         "Blend trees & parametric locomotion (speed-driven)",
         "Cascaded shadow maps for large scenes"],
        size=12.5, gap=6)
page_num(s, 13)


import os
out = os.path.join(os.path.dirname(__file__), "FinalPresentation.pptx")
prs.save(out)
print("Saved:", out, "| slides:", len(prs.slides._sldIdLst))
